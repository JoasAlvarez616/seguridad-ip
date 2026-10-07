# ============================================================
#  SEGURIDAD IP — Diapositivas estilo editorial + interactivas
#  Versión Mejorada: Sin superposiciones, con información completa
# ============================================================
import math
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE

# ---------------- DESIGN TOKENS ----------------
INK      = RGBColor(0x0B, 0x0F, 0x14)   # casi negro
PAPER    = RGBColor(0xF4, 0xF1, 0xEA)   # crema
PAPER_2  = RGBColor(0xEA, 0xE4, 0xD6)   # crema más oscuro
VERM     = RGBColor(0xE8, 0x3B, 0x1E)   # bermellón
TEAL     = RGBColor(0x18, 0x7B, 0x7B)   # verde azulado
GOLD     = RGBColor(0xB8, 0x8A, 0x2E)   # ocre
MUTED_L  = RGBColor(0x8A, 0x88, 0x80)   # gris sobre claro
MUTED_D  = RGBColor(0x5E, 0x62, 0x68)   # gris sobre oscuro
WHITE    = RGBColor(0xFF, 0xFF, 0xFF)
GHOST_D  = RGBColor(0x1A, 0x20, 0x28)   # fantasma sobre oscuro
GHOST_L  = RGBColor(0xDD, 0xD6, 0xC5)   # fantasma sobre claro

SERIF = "Georgia"
SANS  = "Arial"

SW, SH = Inches(13.333), Inches(7.5)
M      = Inches(0.85)            # margen lateral
CW     = SW - M * 2              # ancho de contenido

# ---------------- HELPERS ----------------
def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])

def bg(slide, color):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
    s.fill.solid(); s.fill.fore_color.rgb = color
    s.line.fill.background(); s.shadow.inherit = False
    return s

def hair(slide, x, y, w, color, weight=0.75):
    s = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x, y, x + w, y)
    s.line.color.rgb = color
    s.line.width = Pt(weight)
    return s

def vhair(slide, x, y, h, color, weight=0.75):
    s = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x, y, x, y + h)
    s.line.color.rgb = color
    s.line.width = Pt(weight)
    return s

def rect(slide, x, y, w, h, fill=None, line=None, lw=0.75):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    if fill is None: s.fill.background()
    else:
        s.fill.solid(); s.fill.fore_color.rgb = fill
    if line is None: s.line.fill.background()
    else:
        s.line.color.rgb = line; s.line.width = Pt(lw)
    s.shadow.inherit = False
    return s

def circ(slide, x, y, d, fill=None, line=None, lw=0.75):
    s = slide.shapes.add_shape(MSO_SHAPE.OVAL, x, y, d, d)
    if fill is None: s.fill.background()
    else:
        s.fill.solid(); s.fill.fore_color.rgb = fill
    if line is None: s.line.fill.background()
    else:
        s.line.color.rgb = line; s.line.width = Pt(lw)
    s.shadow.inherit = False
    return s

def tri(slide, x, y, w, h, fill=None, line=None, lw=0.75, rot=180):
    s = slide.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE, x, y, w, h)
    if fill is None: s.fill.background()
    else:
        s.fill.solid(); s.fill.fore_color.rgb = fill
    if line is None: s.line.fill.background()
    else:
        s.line.color.rgb = line; s.line.width = Pt(lw)
    s.rotation = rot
    s.shadow.inherit = False
    return s

def txt(slide, text, x, y, w, h,
        font=SANS, size=12, color=INK, bold=False, italic=False,
        align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
        spacing=1.15, space_after=0):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame; tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    for m in ("margin_left","margin_right","margin_top","margin_bottom"):
        setattr(tf, m, 0)
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = spacing
    if space_after: p.space_after = Pt(space_after)
    r = p.add_run(); r.text = text
    r.font.name = font; r.font.size = Pt(size)
    r.font.bold = bold; r.font.italic = italic
    r.font.color.rgb = color
    return tb

def multiline(slide, lines, x, y, w, h, font=SANS, size=12, color=INK,
              spacing=1.35, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame; tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    for m in ("margin_left","margin_right","margin_top","margin_bottom"):
        setattr(tf, m, 0)
    tf.vertical_anchor = anchor
    for i, (t, b, it) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.line_spacing = spacing
        r = p.add_run(); r.text = t
        r.font.name = font; r.font.size = Pt(size)
        r.font.bold = b; r.font.italic = it
        r.font.color.rgb = color
    return tb

def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text

def link(slide, shape, target_slide):
    shape.click_action.target_slide = target_slide

def eyebrow(slide, x, y, text, color=MUTED_L, size=10):
    spaced = "  ".join(list(text.upper()))
    txt(slide, spaced, x, y, Inches(6), Inches(0.3),
        font=SANS, size=size, color=color, bold=True)

def btn_volver(slide, color_line, color_text):
    """Crea el botón interactivo y lo marca con un nombre para identificarlo."""
    b = rect(slide, SW - M - Inches(2.3), Inches(0.55), Inches(2.3), Inches(0.45),
             fill=None, line=color_line, lw=0.75)
    b.name = "btn_volver"
    t = txt(slide, "←  Volver al índice", SW - M - Inches(2.3), Inches(0.55),
            Inches(2.3), Inches(0.45), font=SANS, size=10, color=color_text,
            align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    t.name = "btn_volver_txt"
    return b

# ============================================================
#  CONSTRUIR
# ============================================================
prs = Presentation()
prs.slide_width, prs.slide_height = SW, SH

idx_holder = {}
div_holder = {}
slides_referencias = {}

# ============================================================
# SLIDE 01 — PORTADA (dark, editorial asimétrico)
# ============================================================
s = blank(prs); bg(s, INK)
rect(s, Inches(-1), Inches(5.9), Inches(15), Inches(0.06), fill=VERM)
circ(s, Inches(11.7), Inches(-1.5), Inches(4.5), line=GHOST_D, lw=0.75)
circ(s, Inches(12.6), Inches(0.6), Inches(2.5), line=GHOST_D, lw=0.75)

eyebrow(s, M, Inches(0.55), "Guía de estudio · Ciberseguridad", MUTED_D, 9)

txt(s, "Seguridad", M, Inches(1.55), Inches(9), Inches(1.5),
    font=SERIF, size=88, color=WHITE, bold=False, spacing=0.95)
txt(s, "IP", M, Inches(2.85), Inches(3), Inches(1.6),
    font=SERIF, size=88, color=VERM, bold=False, italic=True, spacing=0.95)
txt(s, "Protocolos, ataques\ny controles en red",
    M + Inches(2.5), Inches(3.35), Inches(6), Inches(1.4),
    font=SERIF, size=22, color=PAPER, italic=True, spacing=1.25)

hair(s, M, Inches(5.95), Inches(5.5), MUTED_D, 0.75)

# Integrantes actualizados
txt(s, "EXPOSICIÓN GRUPAL", M, Inches(6.15), Inches(5), Inches(0.3),
    font=SANS, size=10, color=MUTED_D, bold=True)

multiline(s,
    [("Joás Alvarez", False, False),
     ("Luis Miranda", False, False),
     ("Jean Amaya", False, False),
     ("Johansen Chapman", False, False)],
    M, Inches(6.45), Inches(5), Inches(0.9),
    font=SANS, size=11, color=PAPER, spacing=1.3)

txt(s, "2026", SW - M - Inches(1.5), Inches(0.55), Inches(1.5), Inches(0.4),
    font=SERIF, size=14, color=MUTED_D, italic=True, align=PP_ALIGN.RIGHT)
notes(s, "Portada. Presentar al grupo (Joás, Luis, Jean, Johansen), anunciar el tema y la estructura.")

# ============================================================
# SLIDE 02 — ÍNDICE INTERACTIVO (papel)
# ============================================================
s = blank(prs); bg(s, PAPER)
slides_referencias["indice"] = s
vhair(s, M, Inches(0.6), Inches(6.3), INK, 0.75)
eyebrow(s, M + Inches(0.25), Inches(0.55), "Índice", MUTED_L, 9)

txt(s, "Cuatro", M + Inches(0.25), Inches(1.1), Inches(5), Inches(1),
    font=SERIF, size=54, color=INK, bold=False)
txt(s, "movimientos", M + Inches(0.25), Inches(2.05), Inches(6), Inches(1),
    font=SERIF, size=54, color=INK, italic=True)

txt(s, "Haz clic en cualquier fila para saltar a la sección.",
    M + Inches(0.25), Inches(3.35), Inches(5), Inches(0.4),
    font=SANS, size=11, color=MUTED_L, italic=True)

index_items = [
    ("01", "Fundamentos y marco normativo", "Qué protegemos y bajo qué normas", VERM),
    ("02", "Protocolos: IPsec, VPN, TLS", "Cómo se cifra y autentica la comunicación", TEAL),
    ("03", "Ataques: spoofing, sniffing, DDoS", "Cómo operan y qué consecuencias generan", VERM),
    ("04", "Protección: segmentación, filtrado", "Controles, mínimo privilegio y cierre", GOLD),
]
row_y = Inches(1.15)
row_h = Inches(1.35)
for i, (num, title, sub, col) in enumerate(index_items):
    y = row_y + i * row_h
    hit = rect(s, M + Inches(6.2), y, Inches(6.0), Inches(1.15))
    txt(s, num, M + Inches(6.2), y - Inches(0.05), Inches(1.3), Inches(1.1),
        font=SERIF, size=44, color=col, italic=True)
    txt(s, title, M + Inches(7.7), y + Inches(0.05), Inches(4.6), Inches(0.5),
        font=SERIF, size=17, color=INK, bold=False)
    txt(s, sub, M + Inches(7.7), y + Inches(0.55), Inches(4.6), Inches(0.4),
        font=SANS, size=11, color=MUTED_L, italic=True)
    txt(s, "→", M + Inches(11.75), y + Inches(0.4), Inches(0.6), Inches(0.5),
        font=SERIF, size=22, color=col, align=PP_ALIGN.RIGHT)
    hair(s, M + Inches(6.2), y + Inches(1.2), Inches(6.0), PAPER_2, 0.75)
    idx_holder[i] = hit

# ============================================================
# SLIDE 03 — DIVISOR 01
# ============================================================
s = blank(prs); bg(s, INK)
div_holder[0] = s
txt(s, "01", Inches(-0.4), Inches(0.4), Inches(8), Inches(6.5), font=SERIF, size=300, color=GHOST_D, italic=True, anchor=MSO_ANCHOR.MIDDLE)
eyebrow(s, M, Inches(0.6), "Sección 01", VERM, 9)
txt(s, "Fundamentos", M, Inches(2.4), Inches(7), Inches(1.2), font=SERIF, size=58, color=WHITE)
txt(s, "y marco normativo", M, Inches(3.4), Inches(7), Inches(1.2), font=SERIF, size=58, color=WHITE, italic=True)
hair(s, M, Inches(4.75), Inches(2.2), VERM, 1.25)
txt(s, "Qué se protege, por qué y bajo qué normas —", M, Inches(5.0), Inches(7.5), Inches(0.5), font=SANS, size=13, color=PAPER)
txt(s, "MSPI · Ley 1581 · ISO 27001 · PHVA", M, Inches(5.35), Inches(7.5), Inches(0.4), font=SANS, size=11, color=MUTED_D, italic=True)
btn_volver(s, MUTED_D, PAPER)

# ============================================================
# SLIDE 04 — ¿QUÉ ES?
# ============================================================
s = blank(prs); bg(s, PAPER)
eyebrow(s, M, Inches(0.55), "01 · Fundamentos", VERM, 9)
hair(s, M, Inches(0.85), CW, PAPER_2, 0.75)
txt(s, "Proteger la información", M, Inches(1.3), Inches(11.5), Inches(1), font=SERIF, size=44, color=INK)
txt(s, "que viaja por redes IP", M, Inches(2.15), Inches(11.5), Inches(1), font=SERIF, size=44, color=VERM, italic=True)
txt(s, "IPv4 · IPv6", M, Inches(3.4), Inches(2), Inches(0.3), font=SANS, size=10, color=MUTED_L, bold=True)
txt(s, "La seguridad IP reúne técnicas, políticas y controles para que la información que viaja por la red no sea leída, alterada ni interrumpida por terceros no autorizados.",
    M, Inches(3.75), Inches(6.6), Inches(1.5), font=SANS, size=13, color=INK, spacing=1.55)
vhair(s, M + Inches(8.2), Inches(3.5), Inches(2.4), PAPER_2, 0.75)
txt(s, "«La seguridad no es un producto que se compra; es un proceso que se gestiona.»",
    M + Inches(8.5), Inches(3.55), Inches(3.1), Inches(2), font=SERIF, size=14, color=TEAL, italic=True, spacing=1.4)
y = Inches(6.3)
hair(s, M, y, CW, PAPER_2, 0.75)
txt(s, "Capa 3–7", M, y + Inches(0.15), Inches(2), Inches(0.35), font=SERIF, size=12, color=INK, italic=True)
txt(s, "IPv4 e IPv6", M + Inches(3.5), y + Inches(0.15), Inches(2.5), Inches(0.35), font=SERIF, size=12, color=INK, italic=True)
txt(s, "RFC · ISO · MinTIC", M + Inches(7.5), y + Inches(0.15), Inches(3), Inches(0.35), font=SERIF, size=12, color=INK, italic=True)

# ============================================================
# SLIDE 05 — 4 PROPIEDADES
# ============================================================
s = blank(prs); bg(s, INK)
eyebrow(s, M, Inches(0.55), "01 · Fundamentos", VERM, 9)
txt(s, "Lo que se protege", M, Inches(1.0), Inches(6), Inches(0.8), font=SERIF, size=36, color=WHITE)
props = [
    ("C", "Confidencialidad", "Solo quien debe leer, lee.", VERM),
    ("I", "Integridad", "Nadie altera el dato en tránsito.", TEAL),
    ("D", "Disponibilidad", "El servicio permanece accesible.", GOLD),
    ("A", "Autenticidad", "Se sabe quién se comunica con quién.", WHITE),
]
start_y = Inches(2.05); step = Inches(1.15)
for i, (letter, name, desc, col) in enumerate(props):
    y = start_y + i * step
    txt(s, letter, M, y - Inches(0.15), Inches(1.1), Inches(1.1), font=SERIF, size=54, color=col, italic=True)
    vhair(s, M + Inches(1.1), y + Inches(0.05), Inches(0.75), GHOST_D, 0.75)
    txt(s, name, M + Inches(1.4), y + Inches(0.02), Inches(4.2), Inches(0.5), font=SERIF, size=22, color=WHITE)
    txt(s, desc, M + Inches(5.8), y + Inches(0.15), Inches(6), Inches(0.5), font=SANS, size=12, color=MUTED_D)
    hair(s, M + Inches(1.4), y + Inches(0.85), Inches(10.4), GHOST_D, 0.5)

# ============================================================
# SLIDE 06 — CADENA DEL RIESGO + PHVA
# ============================================================
s = blank(prs); bg(s, PAPER)
eyebrow(s, M, Inches(0.55), "01 · Fundamentos", VERM, 9)
txt(s, "La cadena del riesgo", M, Inches(1.05), Inches(8), Inches(0.8), font=SERIF, size=34, color=INK)
chain = [("Activo", "Lo que tiene valor", VERM),
         ("Amenaza", "Quién o qué daña", INK),
         ("Vulnerabilidad", "Debilidad explotable", INK),
         ("Riesgo", "Probabilidad × impacto", INK),
         ("Incidente", "Compromiso real", VERM)]
n = len(chain); gap = Inches(0.35); node_w = (CW - (n - 1) * gap) / n
y = Inches(2.15)
for i, (name, desc, col) in enumerate(chain):
    x = M + i * (node_w + gap)
    txt(s, f"{i+1:02d}", x, y, Inches(0.6), Inches(0.35), font=SERIF, size=13, color=MUTED_L, italic=True)
    hair(s, x, y + Inches(0.45), node_w, col, 1.0)
    txt(s, name, x, y + Inches(0.55), node_w, Inches(0.4), font=SERIF, size=15, color=INK, bold=False)
    txt(s, desc, x, y + Inches(1.0), node_w, Inches(0.7), font=SANS, size=10, color=MUTED_L, spacing=1.3)
    if i < n - 1:
        txt(s, "→", x + node_w - Inches(0.05), y + Inches(0.35), gap, Inches(0.3), font=SERIF, size=14, color=MUTED_L, align=PP_ALIGN.CENTER)
y = Inches(5.1)
hair(s, M, y, CW, PAPER_2, 0.75)
txt(s, "Ciclo PHVA", M, y + Inches(0.2), Inches(3), Inches(0.5), font=SERIF, size=22, color=TEAL, italic=True)
txt(s, "Enfoque del MSPI colombiano — la seguridad se gestiona en ciclo continuo.", M, y + Inches(0.75), Inches(7), Inches(0.5), font=SANS, size=11, color=MUTED_L, italic=True)
phva = [("P", "Planear"), ("H", "Hacer"), ("V", "Verificar"), ("A", "Actuar")]
x0 = M + Inches(8.2)
for i, (l, name) in enumerate(phva):
    x = x0 + i * Inches(1.05)
    circ(s, x, y + Inches(0.3), Inches(0.75), line=TEAL, lw=0.75)
    txt(s, l, x, y + Inches(0.3), Inches(0.75), Inches(0.75), font=SERIF, size=22, color=TEAL, italic=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    txt(s, name, x - Inches(0.15), y + Inches(1.2), Inches(1.05), Inches(0.4), font=SANS, size=9, color=MUTED_L, align=PP_ALIGN.CENTER)

# ============================================================
# SLIDE 07 — MARCO NORMATIVO
# ============================================================
s = blank(prs); bg(s, INK)
eyebrow(s, M, Inches(0.55), "01 · Fundamentos", VERM, 9)
txt(s, "Marco normativo", M, Inches(1.05), Inches(8), Inches(0.8), font=SERIF, size=36, color=WHITE)
norms = [
    ("2021", "Resolución 500 de 2021", "MinTIC — lineamientos del MSPI: riesgos, controles e incidentes.", VERM),
    ("2012", "Ley 1581 de 2012", "Protección de datos personales en Colombia.", WHITE),
    ("2022", "Decreto 338 de 2022", "Gobernanza, riesgo e incidentes en entidades públicas.", WHITE),
    ("2022", "ISO/IEC 27001:2022", "Sistema de Gestión de Seguridad de la Información (SGSI).", WHITE),
    ("2009", "Ley 1273 de 2009", "Tipifica delitos informáticos en Colombia.", VERM),
]
y = Inches(2.05); row_h = Inches(0.92)
vhair(s, M + Inches(1.15), y, row_h * len(norms), MUTED_D, 0.5)
for i, (year, name, desc, col) in enumerate(norms):
    ry = y + i * row_h
    txt(s, year, M, ry, Inches(1.0), Inches(0.5), font=SERIF, size=15, color=col, italic=True)
    circ(s, M + Inches(1.05), ry + Inches(0.08), Inches(0.22), fill=col)
    txt(s, name, M + Inches(1.55), ry - Inches(0.02), Inches(4), Inches(0.45), font=SERIF, size=17, color=WHITE)
    txt(s, desc, M + Inches(5.8), ry + Inches(0.05), Inches(6), Inches(0.6), font=SANS, size=11, color=MUTED_D, spacing=1.4)

# ============================================================
# SLIDE 08 — DIVISOR 02
# ============================================================
s = blank(prs); bg(s, INK)
div_holder[1] = s
txt(s, "02", SW - Inches(7.5), Inches(0.4), Inches(8), Inches(6.5), font=SERIF, size=300, color=GHOST_D, italic=True, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
txt(s, "Protocolos", SW - M - Inches(7), Inches(2.4), Inches(7), Inches(1.2), font=SERIF, size=58, color=WHITE, align=PP_ALIGN.RIGHT)
txt(s, "IPsec · VPN · TLS", SW - M - Inches(7), Inches(3.4), Inches(7), Inches(1.2), font=SERIF, size=58, color=TEAL, italic=True, align=PP_ALIGN.RIGHT)
hair(s, SW - M - Inches(2.2), Inches(4.75), Inches(2.2), TEAL, 1.25)
txt(s, "— Cómo se cifra y autentica la comunicación entre extremos", SW - M - Inches(9), Inches(5.0), Inches(9), Inches(0.5), font=SANS, size=13, color=PAPER, align=PP_ALIGN.RIGHT)
txt(s, "RFC 4301 · RFC 8446 · IKEv2", SW - M - Inches(9), Inches(5.35), Inches(9), Inches(0.4), font=SANS, size=11, color=MUTED_D, italic=True, align=PP_ALIGN.RIGHT)
eyebrow(s, SW - M - Inches(3), Inches(0.6), "Sección 02", TEAL, 9)
b = rect(s, M, Inches(0.55), Inches(2.3), Inches(0.45), fill=None, line=MUTED_D, lw=0.75)
b.name = "btn_volver"
txt(s, "←  Volver al índice", M, Inches(0.55), Inches(2.3), Inches(0.45), font=SANS, size=10, color=PAPER, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

# ============================================================
# SLIDE 09 — IPSEC
# ============================================================
s = blank(prs); bg(s, PAPER)
eyebrow(s, M, Inches(0.55), "02 · Protocolos", TEAL, 9)
txt(s, "IPsec", M, Inches(1.05), Inches(5), Inches(0.9), font=SERIF, size=44, color=INK)
txt(s, "Internet Protocol Security", M, Inches(1.9), Inches(5), Inches(0.5), font=SERIF, size=14, color=TEAL, italic=True)
txt(s, "RFC 4301 · IETF", M + Inches(4.5), Inches(1.85), Inches(4), Inches(0.4), font=SANS, size=10, color=MUTED_L)
dy = Inches(2.9)
txt(s, "Paquete original", M, dy - Inches(0.35), Inches(5), Inches(0.3), font=SANS, size=10, color=MUTED_L, bold=True)
rect(s, M, dy, Inches(3.2), Inches(0.7), fill=None, line=INK, lw=0.75)
txt(s, "IP  |  Datos", M, dy, Inches(3.2), Inches(0.7), font=SANS, size=12, color=INK, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
txt(s, "→", M + Inches(3.4), dy - Inches(0.05), Inches(0.9), Inches(0.7), font=SERIF, size=32, color=TEAL, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
rx = M + Inches(4.4)
txt(s, "Paquete protegido por IPsec", rx, dy - Inches(0.35), Inches(7.5), Inches(0.3), font=SANS, size=10, color=MUTED_L, bold=True)
rect(s, rx, dy, Inches(1.5), Inches(0.7), fill=None, line=INK, lw=0.75)
txt(s, "IP nuevo", rx, dy, Inches(1.5), Inches(0.7), font=SANS, size=11, color=INK, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
rect(s, rx + Inches(1.55), dy, Inches(1.5), Inches(0.7), fill=TEAL)
txt(s, "ESP hdr", rx + Inches(1.55), dy, Inches(1.5), Inches(0.7), font=SANS, size=11, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
rect(s, rx + Inches(3.1), dy, Inches(2.6), Inches(0.7), fill=None, line=INK, lw=0.75)
txt(s, "Datos cifrados", rx + Inches(3.1), dy, Inches(2.6), Inches(0.7), font=SANS, size=11, color=INK, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
rect(s, rx + Inches(5.75), dy, Inches(1.4), Inches(0.7), fill=TEAL)
txt(s, "ESP auth", rx + Inches(5.75), dy, Inches(1.4), Inches(0.7), font=SANS, size=11, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
y = Inches(4.6)
hair(s, M, y, CW, PAPER_2, 0.75)
comps = [
    ("AH", "Autentica e integra. No cifra."),
    ("ESP", "Cifra, autentica e integra. El más usado."),
    ("IKEv2", "Negocia claves y parámetros de seguridad."),
    ("Modo túnel", "Encapsula el paquete IP completo. VPN sitio a sitio."),
]
col_w = CW / 4
for i, (name, desc) in enumerate(comps):
    x = M + i * col_w
    txt(s, name, x, y + Inches(0.2), col_w - Inches(0.3), Inches(0.4), font=SERIF, size=15, color=INK, bold=False)
    txt(s, desc, x, y + Inches(0.65), col_w - Inches(0.3), Inches(0.9), font=SANS, size=10, color=MUTED_L, spacing=1.35)

# ============================================================
# SLIDE 10 — VPN
# ============================================================
s = blank(prs); bg(s, INK)
eyebrow(s, M, Inches(0.55), "02 · Protocolos", TEAL, 9)
txt(s, "VPN", M, Inches(1.05), Inches(4), Inches(0.9), font=SERIF, size=44, color=WHITE)
txt(s, "Red Privada Virtual", M, Inches(1.85), Inches(4), Inches(0.5), font=SERIF, size=14, color=TEAL, italic=True)
vhair(s, M, Inches(2.55), Inches(3.2), MUTED_D, 0.5)
txt(s, "Un canal lógico privado\nsobre una red pública.", M, Inches(2.75), Inches(5), Inches(1), font=SANS, size=13, color=PAPER, spacing=1.45)
txt(s, "No es IPsec. IPsec es un protocolo;\nVPN es la solución que puede\nconstruirse con IPsec, TLS u otras\ntecnologías.",
    M, Inches(3.9), Inches(5), Inches(1.6), font=SANS, size=11, color=MUTED_D, spacing=1.55, italic=True)
rx = M + Inches(6.8)
txt(s, "Tipos", rx, Inches(1.15), Inches(5), Inches(0.5), font=SERIF, size=22, color=WHITE, italic=True)
hair(s, rx, Inches(1.65), Inches(5.5), MUTED_D, 0.5)
types = [
    ("01", "Acceso remoto", "Un usuario a la red institucional.", TEAL),
    ("02", "Sitio a sitio", "Dos redes completas conectadas.", TEAL),
    ("03", "Basada en IPsec", "Opera en capa 3.", TEAL),
    ("04", "Basada en TLS", "Opera en capas superiores.", TEAL),
]
ry = Inches(2.05); rh = Inches(1.05)
for i, (num, name, desc, col) in enumerate(types):
    y = ry + i * rh
    txt(s, num, rx, y, Inches(0.6), Inches(0.4), font=SERIF, size=13, color=col, italic=True)
    txt(s, name, rx + Inches(0.75), y - Inches(0.05), Inches(3.5), Inches(0.45), font=SERIF, size=18, color=WHITE)
    txt(s, desc, rx + Inches(0.75), y + Inches(0.4), Inches(4.5), Inches(0.5), font=SANS, size=11, color=MUTED_D)
    if i < 3: hair(s, rx, y + Inches(0.9), Inches(5.5), GHOST_D, 0.5)

# ============================================================
# SLIDE 11 — TLS vs IPSEC
# ============================================================
s = blank(prs); bg(s, PAPER)
eyebrow(s, M, Inches(0.55), "02 · Protocolos", TEAL, 9)
txt(s, "TLS, el protocolo detrás de HTTPS", M, Inches(1.05), Inches(11.5), Inches(0.9), font=SERIF, size=32, color=INK)
txt(s, "Transport Layer Security · capa 4–7 · RFC 8446 · TLS 1.3 recomendado", M, Inches(1.85), Inches(11.5), Inches(0.4), font=SANS, size=11, color=MUTED_L, italic=True)
y = Inches(2.65)
vhair(s, M + Inches(5.65), y, Inches(4.0), PAPER_2, 0.75)
txt(s, "IPsec", M, y, Inches(5), Inches(0.5), font=SERIF, size=24, color=INK)
txt(s, "capa 3 · red", M, y + Inches(0.5), Inches(5), Inches(0.35), font=SANS, size=11, color=MUTED_L, italic=True)
txt(s, "TLS", M + Inches(6.0), y, Inches(5), Inches(0.5), font=SERIF, size=24, color=TEAL)
txt(s, "capa 4–7 · aplicación", M + Inches(6.0), y + Inches(0.5), Inches(5), Inches(0.35), font=SANS, size=11, color=MUTED_L, italic=True)
rows = [
    ("Qué protege", "Cualquier tráfico IP", "Conexiones de aplicaciones"),
    ("Cifrado", "ESP", "Nativo"),
    ("Uso típico", "VPN sitio a sitio", "HTTPS, correo, web"),
    ("Autenticación", "IKEv2", "Certificados digitales"),
]
ry = y + Inches(1.05)
for label, a, b in rows:
    txt(s, label.upper(), M, ry, Inches(2), Inches(0.3), font=SANS, size=9, color=MUTED_L, bold=True)
    txt(s, a, M, ry + Inches(0.3), Inches(5.2), Inches(0.5), font=SERIF, size=14, color=INK)
    txt(s, b, M + Inches(6.0), ry + Inches(0.3), Inches(5.5), Inches(0.5), font=SERIF, size=14, color=INK)
    hair(s, M, ry + Inches(0.75), Inches(11.5), PAPER_2, 0.5)
    ry += Inches(0.78)

# ============================================================
# SLIDE 12 — DIVISOR 03
# ============================================================
s = blank(prs); bg(s, INK)
div_holder[2] = s
tri(s, Inches(9.5), Inches(-1), Inches(6), Inches(6), fill=GHOST_D, rot=0)
tri(s, Inches(11), Inches(3), Inches(4), Inches(4), fill=GHOST_D, rot=0)
txt(s, "03", Inches(-0.4), Inches(0.4), Inches(8), Inches(6.5), font=SERIF, size=300, color=GHOST_D, italic=True, anchor=MSO_ANCHOR.MIDDLE)
eyebrow(s, M, Inches(0.6), "Sección 03", VERM, 9)
txt(s, "Ataques", M, Inches(2.4), Inches(7), Inches(1.2), font=SERIF, size=58, color=WHITE)
txt(s, "en red", M, Inches(3.4), Inches(7), Inches(1.2), font=SERIF, size=58, color=VERM, italic=True)
hair(s, M, Inches(4.75), Inches(2.2), VERM, 1.25)
txt(s, "— Spoofing · Sniffing · DDoS", M, Inches(5.0), Inches(7.5), Inches(0.5), font=SANS, size=13, color=PAPER)
txt(s, "Cómo operan y qué consecuencias generan", M, Inches(5.35), Inches(7.5), Inches(0.4), font=SANS, size=11, color=MUTED_D, italic=True)
btn_volver(s, MUTED_D, PAPER)

# ============================================================
# SLIDE 13 — SPOOFING (Corregido Overlaps)
# ============================================================
s = blank(prs); bg(s, PAPER)
eyebrow(s, M, Inches(0.55), "03 · Ataques", VERM, 9)
txt(s, "Spoofing", M, Inches(1.05), Inches(6), Inches(0.9), font=SERIF, size=44, color=INK)
txt(s, "Suplantar una identidad o un dato para aparentar ser otro.", M, Inches(2.0), Inches(11.5), Inches(0.5), font=SERIF, size=16, color=VERM, italic=True)

# Coordenadas ajustadas para evitar superposiciones
y = Inches(2.9)
rect(s, M, y, Inches(5.4), Inches(2.7), fill=None, line=INK, lw=0.75)
txt(s, "01", M + Inches(0.4), y + Inches(0.25), Inches(1), Inches(0.5), font=SERIF, size=16, color=VERM, italic=True)
txt(s, "IP spoofing", M + Inches(0.4), y + Inches(0.7), Inches(4.6), Inches(0.6), font=SERIF, size=24, color=INK)
txt(s, "Falsifica la dirección IP de origen. Base de ataques de amplificación y evasión de controles.",
    M + Inches(0.4), y + Inches(1.4), Inches(4.6), Inches(1.0), font=SANS, size=11, color=MUTED_L, spacing=1.5)

rx = M + Inches(5.85)
bw = (CW - Inches(5.85) - Inches(0.4)) / 3
for i, (num, name, desc) in enumerate([
    ("02", "ARP", "IP ↔ MAC en red local."),
    ("03", "DNS", "Respuesta falsa, redirección."),
    ("04", "Email", "Remitente falsificado."),
]):
    x = rx + i * (bw + Inches(0.2))
    hair(s, x, y + Inches(0.1), bw, VERM, 1.0)
    txt(s, num, x, y + Inches(0.25), bw, Inches(0.4), font=SERIF, size=13, color=VERM, italic=True)
    txt(s, name, x, y + Inches(0.7), bw, Inches(0.5), font=SERIF, size=20, color=INK)
    txt(s, desc, x, y + Inches(1.3), bw, Inches(1.2), font=SANS, size=11, color=MUTED_L, spacing=1.4)

y2 = Inches(6.1)
hair(s, M, y2, CW, PAPER_2, 0.75)
txt(s, "MITIGACIÓN", M, y2 + Inches(0.15), Inches(2), Inches(0.3), font=SANS, size=9, color=GOLD, bold=True)
txt(s, "Validación de origen · ACL · filtrado ingreso/salida · uRPF · autenticación fuerte · Dynamic ARP Inspection.",
    M + Inches(1.6), y2 + Inches(0.15), Inches(10), Inches(0.6), font=SANS, size=11, color=INK, spacing=1.35)

# ============================================================
# SLIDE 14 — SNIFFING (Corregido Overlaps)
# ============================================================
s = blank(prs); bg(s, INK)
eyebrow(s, M, Inches(0.55), "03 · Ataques", VERM, 9)
txt(s, "Sniffing", M, Inches(1.05), Inches(6), Inches(0.9), font=SERIF, size=44, color=WHITE)
txt(s, "Capturar y analizar tráfico. No siempre es ataque: el administrador lo hace con autorización. Es ataque cuando se intercepta sin permiso.",
    M, Inches(2.0), Inches(8.5), Inches(0.8), font=SANS, size=12, color=MUTED_D, spacing=1.55)

vy = Inches(3.2)
vhair(s, M, vy + Inches(0.75), Inches(8.5), MUTED_D, 0.5)
for i in range(7):
    x = M + Inches(0.3) + i * Inches(1.15)
    opacity_col = TEAL if i in (2, 4) else GHOST_D
    rect(s, x, vy + Inches(0.55), Inches(0.65), Inches(0.4), fill=opacity_col)
    txt(s, "pkt", x, vy + Inches(0.55), Inches(0.65), Inches(0.4), font=SANS, size=8, color=WHITE if i in (2,4) else MUTED_D, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

txt(s, "capturado por atacante →", M + Inches(2.2), vy + Inches(1.25), Inches(4), Inches(0.3), font=SANS, size=9, color=VERM, italic=True)
txt(s, "▲", M + Inches(2.85), vy + Inches(0.15), Inches(0.4), Inches(0.4), font=SERIF, size=14, color=VERM, align=PP_ALIGN.CENTER)

rx = M + Inches(9.2)
txt(s, "Qué se captura", rx, vy - Inches(0.35), Inches(4), Inches(0.4), font=SERIF, size=16, color=WHITE, italic=True)
hair(s, rx, vy + Inches(0.05), Inches(3.3), MUTED_D, 0.5)
items = ["Contraseñas sin cifrar", "Cookies y tokens", "Correos sin proteger", "Consultas DNS", "Metadatos"]
for i, it in enumerate(items):
    txt(s, "—", rx, vy + Inches(0.25 + i * 0.45), Inches(0.3), Inches(0.3), font=SANS, size=11, color=VERM)
    txt(s, it, rx + Inches(0.35), vy + Inches(0.25 + i * 0.45), Inches(3.2), Inches(0.3), font=SANS, size=11, color=PAPER)

y2 = Inches(6.2)
hair(s, M, y2, CW, MUTED_D, 0.5)
txt(s, "HTTPS · SSH en vez de Telnet · SFTP en vez de FTP · VPN en redes no confiables · segmentación · MFA.",
    M, y2 + Inches(0.15), CW, Inches(0.4), font=SANS, size=11, color=MUTED_D, italic=True)

# ============================================================
# SLIDE 15 — DDOS (Corregido Overlaps)
# ============================================================
s = blank(prs); bg(s, PAPER)
eyebrow(s, M, Inches(0.55), "03 · Ataques", VERM, 9)
txt(s, "DDoS", M, Inches(1.05), Inches(6), Inches(0.9), font=SERIF, size=44, color=INK)
txt(s, "Distributed Denial of Service — la disponibilidad como objetivo.", M, Inches(2.0), Inches(11.5), Inches(0.5), font=SERIF, size=16, color=VERM, italic=True)

cx = Inches(3.4); cy = Inches(4.2)
rect(s, cx, cy, Inches(1.3), Inches(0.9), fill=INK)
txt(s, "servidor", cx, cy, Inches(1.3), Inches(0.9), font=SANS, size=10, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
for i in range(14):
    angle = math.radians(i * (360/14))
    r = Inches(1.6)
    px = cx + Inches(0.65) + int(r * math.cos(angle)) - Inches(0.1)
    py = cy + Inches(0.45) + int(r * math.sin(angle)) - Inches(0.1)
    tri(s, px, py, Inches(0.28), Inches(0.28), fill=VERM, rot=0)

rx = M + Inches(5.8)
txt(s, "Tres vectores", rx, Inches(2.1), Inches(6.5), Inches(0.5), font=SERIF, size=22, color=INK, italic=True)
hair(s, rx, Inches(2.6), Inches(6.5), PAPER_2, 0.75)
types = [
    ("Volumétrico", "Satura el ancho de banda con grandes volúmenes de tráfico."),
    ("De protocolo", "Agota recursos de red o servidor con conexiones incompletas."),
    ("Capa de aplicación", "Sobrecarga la app con solicitudes aparentemente legítimas."),
]
ry = Inches(2.85)
for name, desc in types:
    txt(s, name, rx, ry, Inches(3), Inches(0.4), font=SERIF, size=17, color=INK)
    txt(s, desc, rx + Inches(3.3), ry + Inches(0.05), Inches(3.4), Inches(0.9), font=SANS, size=11, color=MUTED_L, spacing=1.45)
    hair(s, rx, ry + Inches(0.85), Inches(6.5), PAPER_2, 0.5)
    ry += Inches(1.0)

y2 = Inches(6.3)
hair(s, M, y2, CW, PAPER_2, 0.75)
txt(s, "MITIGACIÓN:", M, y2 + Inches(0.15), Inches(1.4), Inches(0.3), font=SANS, size=9, color=GOLD, bold=True)
txt(s, "anti-DDoS · CDN · rate limiting · firewall · IPS · monitoreo · plan de respuesta · coordinación con el ISP.",
    M + Inches(1.3), y2 + Inches(0.15), Inches(11), Inches(0.4), font=SANS, size=11, color=INK)

# ============================================================
# SLIDE 16 — DIVISOR 04
# ============================================================
s = blank(prs); bg(s, INK)
div_holder[3] = s
txt(s, "04", Inches(-0.4), Inches(0.4), Inches(8), Inches(6.5), font=SERIF, size=300, color=GHOST_D, italic=True, anchor=MSO_ANCHOR.MIDDLE)
eyebrow(s, M, Inches(0.6), "Sección 04", GOLD, 9)
txt(s, "Protección", M, Inches(2.4), Inches(7), Inches(1.2), font=SERIF, size=58, color=WHITE)
txt(s, "y respuesta", M, Inches(3.4), Inches(7), Inches(1.2), font=SERIF, size=58, color=GOLD, italic=True)
hair(s, M, Inches(4.75), Inches(2.2), GOLD, 1.25)
txt(s, "— Segmentación · filtrado · mínimo privilegio", M, Inches(5.0), Inches(7.5), Inches(0.5), font=SANS, size=13, color=PAPER)
txt(s, "Controles que reducen el riesgo y limitan el impacto.", M, Inches(5.35), Inches(7.5), Inches(0.4), font=SANS, size=11, color=MUTED_D, italic=True)
btn_volver(s, MUTED_D, PAPER)

# ============================================================
# SLIDE 17 — SEGMENTACIÓN
# ============================================================
s = blank(prs); bg(s, PAPER)
eyebrow(s, M, Inches(0.55), "04 · Protección", GOLD, 9)
txt(s, "Segmentación de redes", M, Inches(1.05), Inches(11), Inches(0.9), font=SERIF, size=34, color=INK)
txt(s, "Dividir la red en zonas con diferente nivel de confianza. Base del modelo Zero Trust: no confiar en nada, verificar todo.",
    M, Inches(1.85), Inches(11.5), Inches(0.6), font=SANS, size=12, color=MUTED_L, spacing=1.5)

y = Inches(3.0)
zones = [
    ("Internet", "público", MUTED_L, GOLD),
    ("DMZ", "portal · correo", WHITE, GOLD),
    ("Servidores", "web · BD", WHITE, INK),
    ("Académica", "estudiantes", WHITE, INK),
    ("Administrativa", "nómina · finanzas", WHITE, INK),
    ("Invitados", "visitantes", WHITE, GOLD),
    ("IoT", "cámaras", WHITE, GOLD),
]
bw = Inches(1.55); gap = Inches(0.15)
for i, (name, sub, fillc, linec) in enumerate(zones):
    x = M + i * (bw + gap)
    rect(s, x, y, bw, Inches(1.3), fill=fillc, line=linec, lw=0.75)
    txt(s, name, x, y + Inches(0.15), bw, Inches(0.4), font=SERIF, size=12, color=INK, align=PP_ALIGN.CENTER)
    txt(s, sub, x, y + Inches(0.6), bw, Inches(0.6), font=SANS, size=9, color=MUTED_L, align=PP_ALIGN.CENTER)
for i in range(len(zones) - 1):
    x = M + (i + 1) * (bw + gap) - Inches(0.14)
    txt(s, "→", x - Inches(0.05), y + Inches(0.55), Inches(0.2), Inches(0.3), font=SERIF, size=12, color=MUTED_L, align=PP_ALIGN.CENTER)

y2 = Inches(5.0)
hair(s, M, y2, CW, PAPER_2, 0.75)
txt(s, "BENEFICIOS", M, y2 + Inches(0.2), Inches(2), Inches(0.3), font=SANS, size=9, color=GOLD, bold=True)
benefits = [
    ("Limita movimiento lateral", "Una intrusión no se propaga."),
    ("Reduce impacto malware", "Contención por segmento."),
    ("Facilita monitoreo", "Tráfico por zona, más visible."),
    ("Mínimo privilegio", "Cada zona con sus reglas."),
]
col_w = CW / 4
for i, (name, desc) in enumerate(benefits):
    x = M + i * col_w
    txt(s, name, x, y2 + Inches(0.55), col_w - Inches(0.3), Inches(0.4), font=SERIF, size=14, color=INK)
    txt(s, desc, x, y2 + Inches(1.0), col_w - Inches(0.3), Inches(0.6), font=SANS, size=10, color=MUTED_L, spacing=1.35)

# ============================================================
# SLIDE 18 — FILTRADO
# ============================================================
s = blank(prs); bg(s, INK)
eyebrow(s, M, Inches(0.55), "04 · Protección", GOLD, 9)
txt(s, "Filtrado de tráfico", M, Inches(1.05), Inches(8), Inches(0.9), font=SERIF, size=36, color=WHITE)
txt(s, "Permitir, bloquear o inspeccionar según reglas: IP, puerto, protocolo, usuario, aplicación, horario, reputación, comportamiento.",
    M, Inches(1.9), Inches(11), Inches(0.6), font=SANS, size=12, color=MUTED_D, spacing=1.5)

fx = M; fy = Inches(3.0)
rect(s, fx, fy, Inches(4.5), Inches(0.7), fill=None, line=MUTED_D, lw=0.75)
txt(s, "Perímetro", fx + Inches(0.3), fy, Inches(2), Inches(0.7), font=SERIF, size=14, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
txt(s, "Firewall · ACL", fx + Inches(2.3), fy, Inches(2), Inches(0.7), font=SANS, size=11, color=MUTED_D, anchor=MSO_ANCHOR.MIDDLE)

fy2 = fy + Inches(1.0)
rect(s, fx, fy2, Inches(4.5), Inches(0.7), fill=None, line=MUTED_D, lw=0.75)
txt(s, "Inspección", fx + Inches(0.3), fy2, Inches(2), Inches(0.7), font=SERIF, size=14, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
txt(s, "IDS · IPS · WAF", fx + Inches(2.3), fy2, Inches(2), Inches(0.7), font=SANS, size=11, color=MUTED_D, anchor=MSO_ANCHOR.MIDDLE)

fy3 = fy2 + Inches(1.0)
rect(s, fx, fy3, Inches(4.5), Inches(0.7), fill=None, line=GOLD, lw=0.75)
txt(s, "Acceso", fx + Inches(0.3), fy3, Inches(2), Inches(0.7), font=SERIF, size=14, color=GOLD, anchor=MSO_ANCHOR.MIDDLE)
txt(s, "NAC · SIEM", fx + Inches(2.3), fy3, Inches(2), Inches(0.7), font=SANS, size=11, color=GOLD, anchor=MSO_ANCHOR.MIDDLE)

for y0 in (fy + Inches(0.75), fy2 + Inches(0.75)):
    txt(s, "↓", fx + Inches(2.1), y0 - Inches(0.05), Inches(0.3), Inches(0.3), font=SERIF, size=14, color=MUTED_D, align=PP_ALIGN.CENTER)

rx = M + Inches(6.2)
txt(s, "Herramientas", rx, fy - Inches(0.5), Inches(5), Inches(0.5), font=SERIF, size=18, color=WHITE, italic=True)
hair(s, rx, fy, Inches(5.5), MUTED_D, 0.5)
tools = [
    ("Firewall", "Controla tráfico entre redes."),
    ("ACL", "Reglas en routers y switches."),
    ("IDS", "Detecta intrusiones (alerta)."),
    ("IPS", "Previene intrusiones (bloquea)."),
    ("WAF", "Protege apps web y APIs."),
    ("NAC", "Controla quién se conecta."),
    ("SIEM", "Centraliza y correlaciona registros."),
]
ty = fy + Inches(0.2)
for name, desc in tools:
    txt(s, name, rx, ty, Inches(1.5), Inches(0.4), font=SERIF, size=13, color=GOLD)
    txt(s, desc, rx + Inches(1.6), ty, Inches(4), Inches(0.4), font=SANS, size=11, color=MUTED_D)
    ty += Inches(0.4)

# ============================================================
# SLIDE 19 — PLAN + MÍNIMO PRIVILEGIO (Corregido Overlaps)
# ============================================================
s = blank(prs); bg(s, PAPER)
eyebrow(s, M, Inches(0.55), "04 · Protección", GOLD, 9)
txt(s, "Plan de protección", M, Inches(1.05), Inches(11), Inches(0.9), font=SERIF, size=34, color=INK)
txt(s, "10", M, Inches(2.15), Inches(4), Inches(3.2), font=SERIF, size=200, color=GHOST_L, italic=True, anchor=MSO_ANCHOR.MIDDLE)
txt(s, "controles", M + Inches(2.4), Inches(5.6), Inches(3), Inches(0.5), font=SERIF, size=22, color=GOLD, italic=True)

rx = M + Inches(4.8)
plan = [
    "Clasificar activos e información crítica.",
    "Segmentar la red por función y criticidad.",
    "Firewall y ACL: denegar por defecto.",
    "VPN + MFA para acceso remoto.",
    "TLS actualizado en servicios web.",
    "IPsec entre sedes.",
    "Equipos y sistemas actualizados.",
    "Monitoreo con IDS, IPS y SIEM.",
    "Plan de respuesta a incidentes.",
    "Ciclo PHVA continuo.",
]
# Altura de las filas ajustadas
ry = Inches(1.9); rh = Inches(0.48)
for i, item in enumerate(plan):
    y = ry + i * rh
    txt(s, f"{i+1:02d}", rx, y, Inches(0.5), Inches(0.4), font=SERIF, size=11, color=GOLD, italic=True)
    txt(s, item, rx + Inches(0.6), y - Inches(0.02), Inches(7.5), Inches(0.4), font=SANS, size=12, color=INK)
    if i < len(plan) - 1:
        hair(s, rx, y + Inches(0.38), Inches(7.7), PAPER_2, 0.5)

txt(s, "Mínimo privilegio", M, Inches(6.7), Inches(4), Inches(0.4), font=SERIF, size=13, color=GOLD, italic=True)
txt(s, "Cada quien, solo lo que necesita.", M, Inches(7.05), Inches(4), Inches(0.4), font=SANS, size=10, color=MUTED_L)

# ============================================================
# SLIDE 20 — CIERRE + GRACIAS
# ============================================================
s = blank(prs); bg(s, INK)
slides_referencias["cierre"] = s
circ(s, Inches(10.5), Inches(-2), Inches(6), line=GHOST_D, lw=0.75)
circ(s, Inches(11.8), Inches(3.5), Inches(4), line=GHOST_D, lw=0.75)
eyebrow(s, M, Inches(0.6), "Cierre", GOLD, 9)
txt(s, "La seguridad IP no depende", M, Inches(1.6), Inches(11), Inches(0.9), font=SERIF, size=36, color=WHITE)
txt(s, "de una sola tecnología.", M, Inches(2.4), Inches(11), Inches(0.9), font=SERIF, size=36, color=GOLD, italic=True)
hair(s, M, Inches(3.6), Inches(3), GOLD, 1.25)
txt(s, "IPsec, VPN y TLS protegen los datos en tránsito. La segmentación limita el alcance de una intrusión. El filtrado controla el tráfico. Y la gestión de riesgos —exigida por el marco colombiano y alineada con ISO 27001— hace que la seguridad sea un proceso continuo.",
    M, Inches(3.95), Inches(11.5), Inches(1.6), font=SANS, size=13, color=PAPER, spacing=1.55)
txt(s, "La meta no es una red invulnerable, sino reducir la probabilidad y el impacto.",
    M, Inches(5.85), Inches(11.5), Inches(0.6), font=SERIF, size=18, color=WHITE, italic=True)

y2 = Inches(6.85)
txt(s, "Gracias · Preguntas", M, y2, Inches(6), Inches(0.4), font=SERIF, size=13, color=GOLD, italic=True)
btn_volver(s, MUTED_D, PAPER)

# ============================================================
# HACER CLICABLES LOS HYPERLINKS DEL ÍNDICE (Robusto)
# ============================================================
# Índice -> divisores
for i in range(4):
    link(slides_referencias["indice"], idx_holder[i], div_holder[i])

# Divisores y Cierre -> índice (usando la propiedad 'name')
for slide in list(div_holder.values()) + [slides_referencias["cierre"]]:
    for shp in slide.shapes:
        if shp.name == "btn_volver":
            link(slide, shp, slides_referencias["indice"])

# ---------------- GUARDAR ----------------
out = "Seguridad_IP_Editorial_V2.pptx"
prs.save(out)
print(f"✓ Generado: {out}")
print("  Diseño: editorial · interactivo · superposiciones corregidas · autores agregados")